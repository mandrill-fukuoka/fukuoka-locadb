"""評価用の地点リスト（tools/eval_sample.py）：読み取りのみ・同じ seed なら同じ地点・未判定を一部入れる"""
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "tools"))
import eval_sample  # noqa: E402
from osmworld import World  # noqa: E402
from test_spatial import LAT, LON, city, put  # noqa: E402


@pytest.fixture
def world(db, tmp_path):
    w = World(); city(w)
    for i in range(30):
        w.point(5000 + i, LON + 0.001 * (i + 1), LAT + 0.004, amenity="place_of_worship", name=f"寺{i}")
    w.run(db, tmp_path)
    ids = [r["id"] for r in db.execute("select id from spots where category = 'shrine_temple' order by id")]
    for i, sid in enumerate(ids):
        db.execute("insert into spot_spatial (spot_id, spatial_version, confidence, open_ratio, building_cover_100, water_dist_m, water_kind, road_class, road_width_m) values (%s,'t',%s,%s,%s,%s,%s,%s,%s)",
                   (sid, "low" if i % 5 == 0 else "ok", 0.5, 0.2, 80 if i % 2 else None, "river" if i % 2 else None, "residential", 4.0))
    db.commit()


def run(tmp_path, monkeypatch, seed="s1", n="12"):
    out = tmp_path / f"{seed}.xlsx"
    monkeypatch.setattr(sys, "argv", ["eval_sample", "--n", n, "--seed", seed, "--out", str(out)])
    assert eval_sample.main() == 0
    ws = load_workbook(out)["評価用"]
    return list(ws.iter_rows(min_row=2, values_only=True)), [c.value for c in ws[1]]


def test_sample_is_deterministic_has_low_and_blank_judgement_columns(db, world, tmp_path, monkeypatch):
    rows, head = run(tmp_path, monkeypatch, "s1")
    again, _ = run(tmp_path, monkeypatch, "s1")
    other, _ = run(tmp_path, monkeypatch, "s2")
    ids = lambda rs: [r[1] for r in rs]
    assert ids(rows) == ids(again) and ids(rows) != ids(other)         # 同じ seed なら同じ・違えば違う
    assert len(rows) == len(set(ids(rows))) and len(rows) >= 5          # 重複なし
    assert any(r[10] == "未判定" for r in rows)                            # 未判定の地点も入る
    j = [head.index(h) for h in ("判定：開放感", "判定：建物", "判定：水辺", "判定：道", "判定：タグ", "メモ")]
    assert all(r[i] is None for r in rows for i in j)                    # 人が埋める列は空
    assert all(str(r[7]).startswith("https://www.google.com/maps") for r in rows)


def test_sample_writes_nothing(db, world, tmp_path, monkeypatch):
    before = db.execute("select (select count(*) from spots) s, (select count(*) from spot_spatial) p, (select count(*) from spot_tags) t").fetchone()
    run(tmp_path, monkeypatch)
    assert db.execute("select (select count(*) from spots) s, (select count(*) from spot_spatial) p, (select count(*) from spot_tags) t").fetchone() == before
