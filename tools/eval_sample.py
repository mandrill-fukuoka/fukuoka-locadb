"""評価用の地点リスト（読み取りのみ）：周辺の特徴の段階とルールのタグが妥当か、人が見て確かめるための一覧（xlsx）

  python tools/eval_sample.py [--n 120] [--seed 2026-10-19] [--out eval_sample.xlsx]

- 優先度A（団地・歴史・寺社・港湾・工場・倉庫・商店街・飲食店街・空き店舗・漁港）を多めに、B・Cを少なめに無作為に選ぶ
  （対象外の種類：坂道・細街路・線路沿い・工事中・河川は、地形や道から作った地点なので入れない）
- 乱数は seed で固定（同じ seed なら同じ地点）。周辺の特徴が「未判定（low）」の地点も一部入れる（判定の妥当性を見るため）
- 人が埋める列：各項目の「合っている／違う」と、タグ・メモ。地点の位置はGoogleマップ・ストリートビューのリンクで見る
  （リンクを開くだけで、画像は保存しない）
"""
import argparse
import os
import sys

import psycopg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from psycopg.rows import dict_row

PRIORITY = {
    "A": ["housing_estate", "heritage", "shrine_temple", "port", "factory", "warehouse", "shotengai",
          "eatery_street", "vacant", "fishing_port"],
    "B": ["park", "school", "bridge_overpass", "station", "new_building", "commercial_facility", "island",
          "beach", "coast", "hill", "pond"],
    "C": ["hospital", "civic"],
}
SHARE = {"A": 0.58, "B": 0.34, "C": 0.08}
LOW_SHARE = 0.10              # 全体のうち、周辺の特徴が未判定の地点の割合
BAND = {"open": {"open": "開けている", "mid": "ふつう", "closed": "囲まれている"},
        "built": {"few": "少ない", "mid": "中くらい", "many": "多い"},
        "water": {"near": "すぐそば", "close": "近い", "none": "150m以内に水辺なし"},
        "road": {"narrow": "狭い道", "normal": "ふつうの道", "wide": "広い道"}}
WATER_KIND = {"river": "川", "coast": "海", "pond": "池", "beach": "砂浜"}

SQL = """
select s.id, s.name, s.category as cat_code, sc.name as category, s.area_code, s.address,
       ST_Y(s.geom::geometry) lat, ST_X(s.geom::geometry) lon, s.summary,
       ss.confidence, ss.open_ratio, ss.building_cover_100, ss.building_count_100, ss.water_dist_m, ss.water_kind,
       ss.road_class, ss.road_width_m, ss.road_width_src,
       sp_open(ss.open_ratio) b_open, sp_built(ss.building_cover_100) b_built, sp_water(ss.water_dist_m) b_water,
       sp_road(ss.road_class, ss.road_width_m) b_road,
       (select string_agg(t.name, '・' order by t.code) from spot_tags st join tags t on t.code = st.tag_code
         where st.spot_id = s.id and st.origin = 'rule') as rule_tags,
       row_number() over (partition by s.category order by md5(s.id::text || %(seed)s)) as rn
  from spots s
  join spot_categories sc on sc.code = s.category
  join spot_spatial ss on ss.spot_id = s.id
 where s.data_status <> 'archived' and s.category = any(%(cats)s) and (ss.confidence = %(conf)s)
"""


def pick(conn, cats: list[str], n: int, seed: str, conf: str) -> list[dict]:
    """種類ごとに均等に近く配る（少ない種類は全部）"""
    if n <= 0 or not cats:
        return []
    rows = conn.execute(SQL, {"seed": seed, "cats": cats, "conf": conf}).fetchall()
    by: dict[str, list] = {}
    for r in rows:
        by.setdefault(r["cat_code"], []).append(r)
    out, k = [], 0
    while len(out) < n and any(by.values()):
        for c in sorted(by):
            if by[c] and len(out) < n:
                out.append(by[c].pop(0))
        k += 1
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--n", type=int, default=120)
    ap.add_argument("--seed", default="2026-10-19")
    ap.add_argument("--out", default="eval_sample.xlsx")
    a = ap.parse_args()
    n_low = round(a.n * LOW_SHARE)
    with psycopg.connect(os.environ["DATABASE_URL"], row_factory=dict_row) as conn:
        conn.execute("set transaction read only")
        chosen = []
        for lv in "ABC":
            chosen += [dict(r, prio=lv) for r in pick(conn, PRIORITY[lv], round((a.n - n_low) * SHARE[lv]), a.seed, "ok")]
        allcats = [c for lv in "ABC" for c in PRIORITY[lv]]
        chosen += [dict(r, prio="低") for r in pick(conn, allcats, n_low, a.seed, "low")]
        total = conn.execute("select count(*) n from spot_spatial").fetchone()["n"]
    wb = Workbook(); ws = wb.active; ws.title = "評価用"
    head = ["No", "地点ID", "地点名", "種別", "優先度", "エリア", "住所", "Googleマップ", "ストリートビュー（リンクのみ）",
            "開放性（0〜1）", "開放感", "建物面積率", "建物の多さ", "水辺までm", "水辺", "道幅m", "道", "ルールのタグ",
            "判定：開放感", "判定：建物", "判定：水辺", "判定：道", "判定：タグ", "メモ"]
    ws.append(head)
    for i, r in enumerate(chosen, 1):
        low = r["confidence"] == "low"
        kind = WATER_KIND.get(r["water_kind"] or "", "")
        water = "" if low else (BAND["water"][r["b_water"]] if r["b_water"] == "none" else f'{kind}が{BAND["water"][r["b_water"]]}')
        ws.append([
            i, r["id"], r["name"], r["category"], r["prio"], r["area_code"], r["address"] or "",
            f'https://www.google.com/maps/search/?api=1&query={r["lat"]:.6f},{r["lon"]:.6f}',
            f'https://www.google.com/maps/@?api=1&map_action=pano&viewpoint={r["lat"]:.6f},{r["lon"]:.6f}',
            None if low else round(r["open_ratio"], 2), "未判定" if low else BAND["open"][r["b_open"]],
            None if low else round(r["building_cover_100"], 2), "" if low else BAND["built"][r["b_built"]],
            None if r["water_dist_m"] is None else round(r["water_dist_m"]), water,
            r["road_width_m"], "" if low or not r["b_road"] else BAND["road"][r["b_road"]],
            r["rule_tags"] or "", None, None, None, None, None, None])
    fill = PatternFill("solid", fgColor="DDE6F0")
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = fill; c.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [5, 8, 26, 14, 7, 11, 28, 14, 14, 9, 11, 9, 10, 9, 18, 8, 10, 22, 11, 11, 11, 11, 11, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "D2"
    ws2 = wb.create_sheet("つかいかた")
    for line in ["【評価用の地点リスト】周辺の特徴（推定）の段階と、ルールのタグが、実際の場所と合っているかを人が見て確かめる。",
                 f"地点数 {len(chosen)}（優先度A・B・Cを無作為に。「低」は周辺の特徴が未判定の地点＝判定の妥当性を見るため）。seed={a.seed}",
                 "各行のリンクで地図・ストリートビューを開き、右の「判定」列に ○（合っている）/ ×（違う）を入れ、気づきをメモに書く。",
                 "「判定：タグ」は、付いているルールのタグ（開けた空間・狭い道・緑・工業構造物・坂）が合っているか。付いていないが付くべきタグはメモへ。",
                 "段階の境界：開放感 開けている≥0.7／ふつう0.35〜0.7／囲まれている<0.35。建物 少ない<0.15／中くらい／多い≥0.30。水辺 すぐそば<50m／近い50〜150m／150m以内になし。道 狭い<5m／ふつう／広い≥6.5m。",
                 "×が多い項目は、境界を直すか、その項目を出さないかを決める材料にする（仕様書10章タスクI）。画像は保存しない。"]:
        ws2.append([line])
    ws2.column_dimensions["A"].width = 140
    wb.save(a.out)
    print(f"{len(chosen)}地点を {a.out} に書きました（空間特徴の全件 {total}）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
