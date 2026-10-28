"""流れ検索の確認（読み取りのみ）：確認用シーン文で build_routes を呼び、上位の流れと点数を表にする

  DATABASE_URL=... python tools/flow_eval.py [--scenes data/flow_eval/scenes.txt] [--json out.json] [--xlsx out.xlsx] [--top 3]

--xlsx：評価用の表（段階3）。シーンごとに上位3件を並べ、人が「判定」列に ○（使える）／×（使えない）と理由を入れる。
地図はGoogleマップのルートへのリンクだけ（画像は保存しない）。

シーン文の解釈は画面と同じ web/app/scene.php（キーワードによる簡易解釈）を php-cli で呼ぶ。
最後に取り消す（rollback）。検索の記録 record_route_search は呼ばない＝おすすめの流れは増えない。
（build_routes は検索ごとの一時表を作るので、読み取り専用のトランザクションでは動かない）
"""
import argparse
import json
import os
import statistics
import subprocess
import time
from pathlib import Path

import psycopg

ROOT = Path(__file__).resolve().parent.parent
PHP = ("require 'web/app/scene.php'; echo json_encode(interpret_by_keywords($argv[1]), JSON_UNESCAPED_UNICODE);")


def interpret(scene: str) -> dict:
    out = subprocess.run(["php", "-r", PHP, scene], cwd=ROOT, capture_output=True, text=True, check=True)
    return json.loads(out.stdout)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--scenes", default=str(ROOT / "data" / "flow_eval" / "scenes.txt"))
    ap.add_argument("--json")
    ap.add_argument("--xlsx")
    ap.add_argument("--top", type=int, default=3)
    a = ap.parse_args()
    scenes = [l.strip() for l in Path(a.scenes).read_text(encoding="utf-8").splitlines()
              if l.strip() and not l.lstrip().startswith("#")]
    results, matches = [], []
    with psycopg.connect(os.environ["DATABASE_URL"]) as c:
        # 画面（service_role）が検索ごとの一時表を作れるか
        tmp = c.execute("""select case when exists (select 1 from pg_roles where rolname = 'service_role')
                                  then has_database_privilege('service_role', current_database(), 'TEMP') end""").fetchone()[0]
        print(f"service_role の一時表の権限：{ {True: 'あり', False: 'なし（！画面の検索が失敗します）', None: '-'}[tmp] }")
        for i, scene in enumerate(scenes, 1):
            q = interpret(scene)
            t0 = time.time()
            rs = c.execute("select build_routes(%s, %s::jsonb, %s, null, null, 8)",
                           [q["sequence"], json.dumps(q["tags"]), q["travel_mode"]]).fetchone()[0] or []
            ms = int((time.time() - t0) * 1000)
            matches += [float(r.get("match") or 0) for r in rs]
            results.append({"no": i, "scene": scene, "query": q, "ms": ms, "routes": rs})
            print(f"\n### {i}. {scene}\n")
            print(f"景観：{' → '.join(q['sequence']) or '（なし）'}／タグ：{', '.join(q['tags']) or '-'}／"
                  f"移動：{q['travel_mode']}／結果 {len(rs)} 件／{ms} ms\n")
            if not rs:
                continue
            print("| 順 | 地点 | 移動 | 適合度 | 内訳 |")
            print("|---:|---|---|---:|---|")
            for k, r in enumerate(rs[: a.top], 1):
                spots = " → ".join(f"{s['name']}（{s['group']}）" for s in r["spots"])
                parts = r.get("score_parts") or {}
                pl = "・".join(f"{key} {float(v):.2f}" for key, v in parts.items()) or "-"
                print(f"| {k} | {spots} | {r.get('travel_mode')} {r.get('total_minutes')}分 | "
                      f"{round(float(r.get('match') or 0) * 100)}% | {pl} |")
        c.rollback()
    if matches:
        qs = statistics.quantiles(matches, n=10) if len(matches) >= 10 else []
        print(f"\n### 適合度の分布（全 {len(matches)} 件）\n")
        print(f"最小 {min(matches):.2f}／中央 {statistics.median(matches):.2f}／最大 {max(matches):.2f}"
              + (f"／10%点 {qs[0]:.2f}／90%点 {qs[-1]:.2f}" if qs else "")
              + f"／0.5以上 {sum(m >= 0.5 for m in matches)} 件")
    if a.xlsx:
        write_xlsx(results, a.xlsx, a.top)
    if a.json:
        Path(a.json).write_text(json.dumps(results, ensure_ascii=False, indent=1, default=str), encoding="utf-8")


def route_url(r: dict) -> str:
    from urllib.parse import urlencode
    pts = [f"{x['lat']},{x['lon']}" for x in r["spots"]]
    q = {"api": 1, "origin": pts[0], "destination": pts[-1],
         "travelmode": "driving" if r.get("travel_mode") == "car" else "walking"}
    if len(pts) > 2:
        q["waypoints"] = "|".join(pts[1:-1])
    return "https://www.google.com/maps/dir/?" + urlencode(q)


def write_xlsx(results: list, path: str, top: int) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "評価"
    head = ["No", "シーン文", "読み取った景観の流れ", "雰囲気", "移動", "順位", "流れ（地点）", "移動の目安", "適合度",
            "内訳", "地図", "判定（○／×）", "理由・メモ"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="44546A")
    dv = DataValidation(type="list", formula1='"○,×,△"', allow_blank=True)
    ws.add_data_validation(dv)
    band = PatternFill("solid", fgColor="F2F2F2")
    for res in results:
        q = res["query"]
        rows = res["routes"][:top] or [None]
        for k, r in enumerate(rows, 1):
            spots = " → ".join(f"{x.get('label') or x['name']}（{x['group']}）" for x in r["spots"]) if r else "（流れが見つからない）"
            parts = "・".join(f"{key} {float(v):.2f}" for key, v in (r.get("score_parts") or {}).items()) if r else ""
            ws.append([res["no"], res["scene"] if k == 1 else "", " → ".join(q["sequence"]) if k == 1 else "",
                       ", ".join(q["tags"]) if k == 1 else "", q["travel_mode"] if k == 1 else "", k if r else "",
                       spots, f"{r['travel_mode']} {r['total_minutes']}分" if r else "",
                       round(float(r.get("match") or 0), 2) if r else "", parts, "" if not r else "ルート", "", ""])
            row = ws.max_row
            if r:
                ws.cell(row, 11).hyperlink = route_url(r)
                ws.cell(row, 11).font = Font(color="0563C1", underline="single")
            dv.add(ws.cell(row, 12))
            if res["no"] % 2 == 0:
                for c in ws[row]:
                    c.fill = band
    for col, w in zip("ABCDEFGHIJKLM", (5, 36, 24, 22, 7, 5, 60, 11, 7, 30, 7, 11, 30)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    note = wb.create_sheet("説明")
    for line in ["流れ検索の評価（段階3）",
                 "判定：上位3件のそれぞれに、そのシーンで使える流れなら ○、使えないなら ×、条件付きなら △。",
                 "理由・メモ：×・△の理由（景観が違う／遠すぎる／同じような場所／実際は撮れない など）。",
                 "地図：Googleマップでルートを開きます（徒歩または車）。移動の目安は直線距離×1.3からの概算です。",
                 "内訳：signal＝条件の当てはまり、typical＝景観らしさ、short＝短さ、named＝名前のある地点、checked＝確認済み（各0〜1）。",
                 "シーン文の読み取りは、画面と同じキーワードによる簡易解釈です。"]:
        note.append([line])
    note.column_dimensions["A"].width = 110
    wb.save(path)


if __name__ == "__main__":
    main()
