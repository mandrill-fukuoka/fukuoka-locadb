#!/usr/bin/env python3
"""承認済みの事実カード候補（xlsx）から、DBの facts_import に渡す JSON を作る。DBには触れない（読み取りと書き出しだけ）。

  python3 build_facts_json.py fact_cards_reviewed.xlsx facts_import.json [web_candidates_review.xlsx]

- 採否=採 の行だけを使う。保留・不採は出さない
- 出典URLのない行は出さず、unresolved に回す（出典のない事実は載せない）
- Web由来の行（取得元 'Web（公式ページ）'）は、web_candidates_review.xlsx の出典URLで補う（地点ID・項目で突き合わせ）
- Wikipedia 由来は記事名とライセンス（CC BY-SA 4.0）を付ける
"""
import json, sys, urllib.parse
from openpyxl import load_workbook

WEB_ITEM = {"scale": "規模", "built": "創設・開設年", "owner": "管理者・所有者", "status": "現況", "designation": "指定・登録"}
WEB_ORG = {"pref.fukuoka.lg.jp": "福岡県", "city.fukuoka.lg.jp": "福岡市", "ur-net.go.jp": "UR都市機構",
           "shintencho.or.jp": "新天町商店街", "kawabatadori.com": "川端通商店街", "fukuokashigyokyo.com": "福岡市漁業協同組合"}

def org(url):
    host = urllib.parse.urlparse(url).netloc.lower()
    for k, v in WEB_ORG.items():
        if host == k or host.endswith("." + k): return v
    return host

def rows(path, sheet=None):
    ws = load_workbook(path, data_only=True)[sheet] if sheet else load_workbook(path, data_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = next(it)
    return [dict(zip(head, r)) for r in it]

def main(src, out, web=None):
    web_url = {}
    if web:
        for r in rows(web):
            if r.get("採否（採/不採/保留）") == "採" and r.get("出典URL"):
                web_url[(int(r["地点ID"]), WEB_ITEM.get(r["種類"]))] = r["出典URL"]
    items, unresolved, held = [], [], 0
    for r in rows(src, "カード候補"):
        if r.get("採否（採/不採/保留）") != "採":
            held += 1; continue
        sid, item = int(r["地点ID"]), r["項目"]
        kind = "wikipedia" if r["取得元"] == "Wikipedia" else "web"
        url = r.get("出典URL") or (web_url.get((sid, item)) if kind == "web" else None)
        if not url:
            unresolved.append({"spot_id": sid, "name": r["地点名"], "item": item, "value": r["値（直して可）"], "reason": "出典URLなし"}); continue
        d = {"spot_id": sid, "item": item, "value": str(r["値（直して可）"]).strip(), "note": r.get("細目・備考"),
             "source_kind": kind, "source_url": url, "quote": r.get("原文（承認済みの文）")}
        if kind == "wikipedia":
            d["source_title"] = (r.get("出典名（記事）") or "") + " - Wikipedia"; d["license"] = "CC BY-SA 4.0"
        else:
            d["source_title"] = org(url) + " 公式ページ"; d["quote"] = None
        if item == "人口" and (r.get("細目・備考") or "").startswith("時点："):
            d["as_of"] = r["細目・備考"].split("：", 1)[1]; d["note"] = None
        elif item == "管理者・所有者" or (r.get("細目・備考") or "").startswith(("要確認", "疑い", "要手入力")):
            d["note"] = None
        items.append(d)
    json.dump({"items": items, "unresolved": unresolved}, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"出力 {len(items)} 件（対象外 {held} 件、出典なし {len(unresolved)} 件）→ {out}")

if __name__ == "__main__":
    if len(sys.argv) < 3: raise SystemExit(__doc__)
    main(*sys.argv[1:4])
